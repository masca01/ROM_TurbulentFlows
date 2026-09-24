"""
Build ONE Excel workbook for the whole Re-sweep (Alpha0 Re = 50, 60, 70, 80 and
the old Re100 dataset).

Sheet "Summary": one row per dataset with the POD truncations used at each energy
level (k_95, k_99), the beta-VAE latent dimension, and the last (converged) error
of every curve.

Then each dataset gets its own sheet that combines the POD and the beta-VAE
convergence results: a small table per curve (POD @95% energy, POD @99% energy,
NN), two note lines stating the error definitions and the mode counts, and a
native scatter chart overlaying all the error curves, plus a second copy of
the chart with log-log axes below it (the view used in the progress report).

Usage:  python3 make_alpha0_excel.py [random|tail|tail5|tail2.5]
    random (default)  results of the random-validation split
                      ../DATA/AUGMENTED/  ->  Alpha0_convergence.xlsx
    tail              results of the tail-validation split (validation = the LAST
                      10% of the record)
                      ../DATA/AUGMENTED_TAILVAL/  ->  Alpha0_convergence_tailval.xlsx
    tail5             validation = the LAST 5% of the record
                      ../DATA/AUGMENTED_TAILVAL5/  ->  Alpha0_convergence_tailval5.xlsx
    tail2.5           validation = the LAST 2.5% of the record
                      ../DATA/AUGMENTED_TAILVAL2.5/  ->  Alpha0_convergence_tailval2.5.xlsx

Reads <aug dir>/<base>_pod95_convergence.mat, _pod99_convergence.mat and
_nn_convergence.mat.  Curves whose .mat is not present yet are simply left out.
"""
import os, sys
import numpy as np
from scipy.io import loadmat
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Series, Reference
from openpyxl.styles import Font, Alignment

from rom import split as cs

SPLIT    = cs.check_mode(sys.argv[1] if len(sys.argv) > 1 else "random")
_AUG_DIR = cs.aug_dir(SPLIT)
OUT      = cs.excel_path(SPLIT)
os.makedirs(os.path.dirname(OUT), exist_ok=True)

RES   = [50, 60, 70, 80, 100]
BASES = {re: f"dataRe{re}Alpha0_2" for re in (50, 60, 70, 80)}
BASES[100] = "Data2PlatesGap1Re100"          # old dataset (1000 snapshots)
POD_LEVELS = [95, 99]                        # POD energy levels to include
HDR   = Font(bold=True)
_VALTXT = cs.describe(SPLIT)                 # e.g. "LAST 2.5% of the record (tail)"

# colours: POD levels in blues, NN in orange
_COL = {"pod95": "1F77B4", "pod99": "17BECF", "nn": "FF7F0E"}


def _load(base, tag):
    f = os.path.join(_AUG_DIR, f"{base}_{tag}_convergence.mat")
    if not os.path.exists(f):
        return None
    S = loadmat(f, simplify_cells=True)
    return {k: np.atleast_1d(S[k]) for k in S if not k.startswith("__")}


def _scalar(d, key, default=None):
    if d is None or key not in d:
        return default
    v = np.ravel(d[key])
    return v[0] if len(v) else default


def _scatter(log=False):
    """Native scatter chart; log=True gives log10 axes on both n and the error
    (the log-log view of the progress report, where the curves are straight-ish)."""
    ch = ScatterChart()
    ch.title = ("POD vs beta-VAE — validation error vs snapshots"
                + ("  (log-log)" if log else ""))
    ch.x_axis.title = "n (snapshots used)"
    ch.y_axis.title = "unexplained energy fraction (validation)"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    if log:
        ch.x_axis.scaling.logBase = 10
        ch.y_axis.scaling.logBase = 10
        ch.x_axis.majorGridlines = None
    ch.style = 13
    ch.width = 22
    ch.height = 13
    return ch


def _series(charts, ws, xcol, ycol, r0, r1, name, color):
    """Add the same (x, y) series to every chart in `charts`."""
    for ch in charts:
        xref = Reference(ws, min_col=xcol, min_row=r0, max_row=r1)
        yref = Reference(ws, min_col=ycol, min_row=r0, max_row=r1)
        s = Series(yref, xref, title=name)
        s.marker.symbol = "circle"
        s.marker.size = 6
        s.marker.graphicalProperties.solidFill = color
        s.marker.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = 18000
        ch.series.append(s)


def _table(ws, col, headers, cols, start_row):
    for c, h in enumerate(headers):
        ws.cell(row=start_row, column=col + c, value=h).font = HDR
    r0 = start_row + 1
    for c, series in enumerate(cols):
        for i, v in enumerate(series):
            ws.cell(row=r0 + i, column=col + c, value=float(v))
    return r0, r0 + len(cols[0]) - 1


def build_sheet(ws, re, pods, nn):
    """pods = {level: dict or None}."""
    lat = _scalar(nn, "latent_dim", "?")
    lat = int(lat) if lat != "?" else lat
    mode_txt = []
    for lv in POD_LEVELS:
        k = _scalar(pods.get(lv), "n_modes")
        if k is not None:
            mode_txt.append(f"{lv}% energy -> top-{int(k)} modes")

    ws["A1"] = (f"Re {re} — POD vs beta-VAE convergence ({BASES[re]}, cap 1500, "
                f"validation = {_VALTXT})")
    ws["A1"].font = HDR
    ws["A2"] = ("POD error = ‖X_V - UUᵀX_V‖F² / ‖X_V‖F² (unexplained energy fraction);  "
                "NN error = 1 - Ek/100 (energy fraction).  "
                "Same kind of quantity — all curves are directly comparable.")
    ws["A3"] = (f"POD truncation:  {';  '.join(mode_txt) if mode_txt else '—'}.   "
                f"beta-VAE latent dimension = {lat}.")
    ws["A2"].alignment = Alignment(wrap_text=False)
    ws["A3"].alignment = Alignment(wrap_text=False)

    charts = [_scatter(), _scatter(log=True)]     # linear axes, then log-log
    col = 1
    top = 5
    for lv in POD_LEVELS:
        pod = pods.get(lv)
        if pod is None:
            continue
        k = int(_scalar(pod, "n_modes", 0))
        r0, r1 = _table(ws, col, [f"POD{lv} n", f"POD{lv} e_mean ({k} modes)"],
                        [pod["n_vals"], pod["e_mean"]], top)
        _series(charts, ws, col, col + 1, r0, r1,
                f"POD {lv}% energy (top-{k} modes)", _COL[f"pod{lv}"])
        ws.column_dimensions[ws.cell(row=top, column=col).column_letter].width = 12
        ws.column_dimensions[ws.cell(row=top, column=col + 1).column_letter].width = 24
        col += 3
    if nn is not None:
        r0, r1 = _table(ws, col, ["NN n", f"NN e_mean (latent {lat})"],
                        [nn["n_vals"], nn["e_mean"]], top)
        _series(charts, ws, col, col + 1, r0, r1,
                f"beta-VAE (1 - Ek/100, latent {lat})", _COL["nn"])
        ws.column_dimensions[ws.cell(row=top, column=col).column_letter].width = 12
        ws.column_dimensions[ws.cell(row=top, column=col + 1).column_letter].width = 24
        col += 3
    anchor_col = ws.cell(row=top, column=col).column_letter
    ws.add_chart(charts[0], f"{anchor_col}{top}")          # linear chart
    ws.add_chart(charts[1], f"{anchor_col}{top + 27}")     # log-log chart below it


def build_summary(ws, found):
    """found = list of (re, pods, nn)."""
    ws["A1"] = (f"Summary — POD truncations, NN latent dimension and converged errors "
                f"(validation = {_VALTXT})")
    ws["A1"].font = HDR
    ws["A2"] = ("Errors are the unexplained energy fraction at the last sweep point "
                "(all snapshots of the training pool).")
    hdr = ["Re", "dataset", "N snapshots", "pool", "val"]
    for lv in POD_LEVELS:
        hdr += [f"POD modes @{lv}%", f"POD{lv} n_last", f"POD{lv} e_last"]
    hdr += ["NN latent", "NN n_last", "NN e_last", "NN Ek_last (%)"]
    for c, h in enumerate(hdr, start=1):
        ws.cell(row=4, column=c, value=h).font = HDR
        ws.column_dimensions[ws.cell(row=4, column=c).column_letter].width = \
            max(10, min(24, len(h) + 2))

    for i, (re, pods, nn) in enumerate(found):
        r = 5 + i
        any_d = nn or next((p for p in pods.values() if p is not None), None)
        n_pool = _scalar(any_d, "n_pool")
        n_val = _scalar(any_d, "n_val")
        row = [re, BASES[re],
               (int(n_pool) + int(n_val)) if n_pool is not None and n_val is not None else None,
               int(n_pool) if n_pool is not None else None,
               int(n_val) if n_val is not None else None]
        for lv in POD_LEVELS:
            p = pods.get(lv)
            if p is None:
                row += [None, None, None]
            else:
                row += [int(_scalar(p, "n_modes", 0)), int(p["n_vals"][-1]),
                        float(p["e_mean"][-1])]
        if nn is None:
            row += [None, None, None, None]
        else:
            e_last = float(nn["e_mean"][-1])
            row += [int(_scalar(nn, "latent_dim", 0)), int(nn["n_vals"][-1]),
                    e_last, 100.0 * (1.0 - e_last)]
        for c, v in enumerate(row, start=1):
            if v is not None:
                ws.cell(row=r, column=c, value=v)


def main():
    print(f"[make_alpha0_excel]  split = {SPLIT}  ({_AUG_DIR} -> {OUT})")
    found = []
    for re in RES:
        base = BASES[re]
        pods = {lv: _load(base, f"pod{lv}") for lv in POD_LEVELS}
        nn = _load(base, "nn")
        if nn is None and all(p is None for p in pods.values()):
            print(f"  Re{re}: no convergence .mat yet — skipped")
            continue
        found.append((re, pods, nn))
        status = ", ".join(f"POD{lv} {'ok' if pods[lv] is not None else '—'}"
                           for lv in POD_LEVELS)
        print(f"  Re{re}: {status}, NN {'ok' if nn is not None else '—'}")

    if not found:
        print("No datasets had convergence results — nothing written.")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    build_summary(ws, found)
    for re, pods, nn in found:
        ws = wb.create_sheet(title=f"Re{re}")
        build_sheet(ws, re, pods, nn)
    wb.save(OUT)
    print(f"\nsaved -> {OUT}  (Summary + {len(found)} dataset sheet(s))")


if __name__ == "__main__":
    main()
