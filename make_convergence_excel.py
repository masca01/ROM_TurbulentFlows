"""
Build an Excel workbook with the POD and NN convergence data + native charts.

Reads the two Re=50 convergence bundles produced by pod_convergence.py and
nn_convergence.py and writes a single .xlsx into ../../convergence/ with:
    - a POD sheet   (n, e_mean, e_std, modes_kept)      + a convergence chart
    - a NN sheet    (n, e_mean, e_std, Ek%)             + error and Ek charts
    - a Comparison sheet overlaying POD vs NN error on one chart

The charts are real Excel scatter charts (editable in Excel), not pictures.
"""

import os, sys
import numpy as np
from scipy.io import loadmat
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Series, Reference
from openpyxl.styles import Font, Alignment

_HERE     = os.path.dirname(os.path.abspath(__file__))
_AUG_DIR  = os.path.normpath(os.path.join(_HERE, "..", "DATA", "AUGMENTED"))
_SPECIAL  = os.path.normpath(os.path.join(_HERE, "..", ".."))       # "597 - Special Topics"
_OUT_DIR  = os.path.join(_SPECIAL, "convergence")
os.makedirs(_OUT_DIR, exist_ok=True)

# dataset base name; override on the command line, e.g.
#   python3 make_convergence_excel.py Data2PlatesGap1Re100
BASE     = sys.argv[1] if len(sys.argv) > 1 else "Data2PlatesGap1Re50"
OUT_XLSX = os.path.join(_OUT_DIR, f"{BASE}_convergence.xlsx")

HDR = Font(bold=True)


def _load(tag):
    S = loadmat(os.path.join(_AUG_DIR, f"{BASE}_{tag}_convergence.mat"),
                simplify_cells=True)
    return {k: np.atleast_1d(S[k]) for k in S if not k.startswith("__")
            } | {k: S[k] for k in ("n_val", "n_pool", "m_reps") if k in S}


def _write_table(ws, headers, cols, start_row=1):
    """Write a column-major table; return (first_data_row, last_data_row)."""
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.font = HDR; cell.alignment = Alignment(horizontal="center")
    r0 = start_row + 1
    for c, col in enumerate(cols, start=1):
        for i, v in enumerate(col):
            ws.cell(row=r0 + i, column=c, value=float(v))
    return r0, r0 + len(cols[0]) - 1


def _scatter(title, x_title, y_title, w=18, h=11):
    ch = ScatterChart()
    ch.title = title
    ch.x_axis.title = x_title
    ch.y_axis.title = y_title
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.style = 13
    ch.width = w; ch.height = h
    return ch


def _add_series(ch, ws, x_col, y_col, r0, r1, name, style="lineMarker"):
    xref = Reference(ws, min_col=x_col, min_row=r0, max_row=r1)
    yref = Reference(ws, min_col=y_col, min_row=r0, max_row=r1)
    s = Series(yref, xref, title=name)
    s.marker.symbol = "circle"; s.marker.size = 6
    if style == "markerOnly":
        s.graphicalProperties.line.noFill = True
    else:
        s.graphicalProperties.line.width = 18000     # EMU ~1.4pt
    ch.series.append(s)
    return s


def main():
    pod = _load("pod")
    nn  = _load("nn")

    wb = Workbook()

    # ---------- POD sheet ----------
    ws = wb.active; ws.title = "POD"
    ws["F1"] = f"POD spatial-mode convergence — {BASE}"; ws["F1"].font = HDR
    nm = int(np.ravel(pod["n_modes"])[0]) if "n_modes" in pod else 0
    trunc = "full basis" if nm in (0,) else f"top-{nm} modes"
    ws["F2"] = (f"centered POD, {trunc}, {int(pod['n_val'])} held-out val "
                f"snapshots, m = {int(pod['m_reps'])} draws/n")
    r0, r1 = _write_table(
        ws,
        ["n (snapshots)", "e_mean (rel. Frob. error)", "e_std", "modes_kept"],
        [pod["n_vals"], pod["e_mean"], pod["e_std"], pod["modes_kept"]],
        start_row=4)
    ch = _scatter("POD convergence: e_n vs n", "n (snapshots to build POD)",
                  "relative validation error  ||X_V - UUᵀX_V||_F / ||X_V||_F")
    _add_series(ch, ws, 1, 2, r0, r1, "POD e_mean")
    ws.add_chart(ch, "F4")
    for col, wdt in zip("ABCD", (14, 24, 12, 12)):
        ws.column_dimensions[col].width = wdt

    # ---------- NN sheet ----------
    ws = wb.create_sheet("NN")
    lat = float(np.ravel(nn["latent_dim"])[0]) if "latent_dim" in nn else float("nan")
    bet = float(np.ravel(nn["beta"])[0]) if "beta" in nn else float("nan")
    ws["G1"] = f"beta-VAE data convergence — {BASE}"; ws["G1"].font = HDR
    ws["G2"] = (f"latent={lat:g}, beta={bet:g}, {int(nn['n_val'])} held-out REAL "
                f"val, m = {int(nn['m_reps'])} draws/n")
    r0, r1 = _write_table(
        ws,
        ["n (snapshots)", "e_mean (1 - Ek/100)", "e_std", "Ek_mean (%)"],
        [nn["n_vals"], nn["e_mean"], nn["e_std"], nn["ek_mean"]],
        start_row=4)
    ch1 = _scatter("beta-VAE convergence: e_n vs n", "n (snapshots to train VAE)",
                   "relative-energy error  e_n = 1 - Ek/100")
    _add_series(ch1, ws, 1, 2, r0, r1, "VAE e_mean")
    ws.add_chart(ch1, "G4")
    ch2 = _scatter("beta-VAE energy captured: Ek vs n", "n (snapshots to train VAE)",
                   "Ek (%) validation energy reconstructed")
    _add_series(ch2, ws, 1, 4, r0, r1, "VAE Ek (%)")
    ws.add_chart(ch2, "G26")
    for col, wdt in zip("ABCD", (14, 22, 12, 14)):
        ws.column_dimensions[col].width = wdt

    # ---------- Comparison sheet ----------
    ws = wb.create_sheet("Comparison")
    ws["A1"] = "POD vs beta-VAE convergence (Re=50)"; ws["A1"].font = HDR
    ws["A2"] = ("Note: POD error = ‖·‖F ratio (√energy);  NN error = 1 - Ek/100 "
                "(energy).  Different definitions — compare trends, not absolute values.")
    ws["A2"].alignment = Alignment(wrap_text=False)
    # small helper tables so the comparison chart is self-contained
    pr0, pr1 = _write_table(ws, ["POD n", "POD e_mean"],
                            [pod["n_vals"], pod["e_mean"]], start_row=4)
    ncol = 4
    for c, h in enumerate(["NN n", "NN e_mean"], start=ncol):
        cell = ws.cell(row=4, column=c, value=h); cell.font = HDR
    nr0 = 5
    for i, (a, b) in enumerate(zip(nn["n_vals"], nn["e_mean"])):
        ws.cell(row=nr0 + i, column=ncol,     value=float(a))
        ws.cell(row=nr0 + i, column=ncol + 1, value=float(b))
    nr1 = nr0 + len(nn["n_vals"]) - 1

    ch = _scatter("POD vs beta-VAE — validation error vs snapshots",
                  "n (snapshots used)", "relative validation error", w=22, h=13)
    _add_series(ch, ws, 1, 2, pr0, pr1, "POD (‖·‖F ratio)")
    _add_series(ch, ws, ncol, ncol + 1, nr0, nr1, "beta-VAE (1 - Ek/100)")
    ws.add_chart(ch, "G4")
    for col, wdt in zip(["A", "B", "D", "E"], (12, 14, 12, 14)):
        ws.column_dimensions[col].width = wdt

    wb.save(OUT_XLSX)
    print(f"saved -> {OUT_XLSX}")
    print(f"  POD: {len(pod['n_vals'])} points  |  NN: {len(nn['n_vals'])} points")


if __name__ == "__main__":
    main()
